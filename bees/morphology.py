"""Spore size and shape measured inside a detection crop."""

import csv
import logging
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import cv2
import numpy as np
from PIL import Image

from bees.analysis_zone import point_inside_analysis_square
from bees.config_loader import ConfigurationError

logger = logging.getLogger(__name__)

PERCENTILES = (5, 25, 50, 75, 95)
CONTOUR_SOURCE = "contour"
RECT_SOURCE = "min_area_rect"
BBOX_SOURCE = "bbox_ellipse"
SUMMARY_METRICS = (
    "major_px",
    "minor_px",
    "area_px",
    "equivalent_diameter_px",
    "axis_ratio",
    "eccentricity",
    "circularity",
)
METRIC_LABELS = {
    "major_px": "Большая ось, пикс.",
    "minor_px": "Малая ось, пикс.",
    "area_px": "Площадь, пикс²",
    "equivalent_diameter_px": "Эквивалентный диаметр, пикс.",
    "axis_ratio": "Отношение осей",
    "eccentricity": "Эксцентриситет",
    "circularity": "Округлость",
    "major_um": "Большая ось, мкм",
    "minor_um": "Малая ось, мкм",
    "equivalent_diameter_um": "Эквивалентный диаметр, мкм",
    "area_um2": "Площадь, мкм²",
}
LEVEL_LABELS = {
    "run": "Запуск",
    "type": "Тип",
    "probe": "Проба",
    "sample": "Образец",
    "photo": "Фото",
}
MONTAGE_PAD = 6
MONTAGE_LABEL_HEIGHT = 34

DETAIL_FIELDS = (
    ("image_path", "Фото"),
    ("type_name", "Тип"),
    ("probe_name", "Проба"),
    ("sample_name", "Образец"),
    ("spore_index", "Номер"),
    ("confidence", "Уверенность"),
    ("center_x", "Центр X"),
    ("center_y", "Центр Y"),
    ("major_px", "Большая ось, пикс."),
    ("minor_px", "Малая ось, пикс."),
    ("area_px", "Площадь, пикс²"),
    ("equivalent_diameter_px", "Эквивалентный диаметр, пикс."),
    ("axis_ratio", "Отношение осей"),
    ("eccentricity", "Эксцентриситет"),
    ("circularity", "Округлость"),
    ("solidity", "Solidity"),
    ("angle_deg", "Угол большой оси, градусы"),
    ("source", "Источник"),
    ("edge_support_ratio", "Поддержка краёв"),
    ("intensity_diff", "Разница яркости"),
    ("bbox_width_px", "Ширина рамки, пикс."),
    ("bbox_height_px", "Высота рамки, пикс."),
    ("crop_clipped", "Кроп обрезан"),
    ("population", "Популяция"),
)
MICRON_FIELDS = (
    ("major_um", "Большая ось, мкм"),
    ("minor_um", "Малая ось, мкм"),
    ("equivalent_diameter_um", "Эквивалентный диаметр, мкм"),
    ("area_um2", "Площадь, мкм²"),
)


@dataclass
class MorphologySettings:
    """Morphology keys read from config.yaml."""

    confidence: float
    padding_ratio: float
    montage_columns: int
    microns_per_pixel: Optional[float]
    canny_threshold1: int
    canny_threshold2: int
    analysis_square_size: int
    count_inside_square: bool = False


@dataclass
class SporeMeasurement:
    """One spore. The crop is kept for the montage and left out of the table."""

    image_path: str
    type_name: str
    probe_name: str
    sample_name: str
    spore_index: int
    confidence: Optional[float]
    center_x: float
    center_y: float
    major_px: float
    minor_px: float
    area_px: float
    equivalent_diameter_px: float
    axis_ratio: float
    eccentricity: float
    circularity: float
    solidity: float
    angle_deg: float
    source: str
    edge_support_ratio: Optional[float]
    intensity_diff: Optional[float]
    bbox_width_px: Optional[float]
    bbox_height_px: Optional[float]
    crop_clipped: bool
    population: str
    crop: Optional[np.ndarray] = field(default=None, repr=False)

    def table_values(self, scale: Optional[float]) -> Dict[str, object]:
        """CSV and Excel fields. Micron columns appear only when a scale is set."""
        values = {
            "image_path": self.image_path,
            "type_name": self.type_name,
            "probe_name": self.probe_name,
            "sample_name": self.sample_name,
            "spore_index": self.spore_index,
            "confidence": self.confidence,
            "center_x": self.center_x,
            "center_y": self.center_y,
            "major_px": self.major_px,
            "minor_px": self.minor_px,
            "area_px": self.area_px,
            "equivalent_diameter_px": self.equivalent_diameter_px,
            "axis_ratio": self.axis_ratio,
            "eccentricity": self.eccentricity,
            "circularity": self.circularity,
            "solidity": self.solidity,
            "angle_deg": self.angle_deg,
            "source": self.source,
            "edge_support_ratio": self.edge_support_ratio,
            "intensity_diff": self.intensity_diff,
            "bbox_width_px": self.bbox_width_px,
            "bbox_height_px": self.bbox_height_px,
            "crop_clipped": self.crop_clipped,
            "population": self.population,
        }
        if scale is not None:
            values.update(micron_values(self, scale))
        return values


@dataclass
class MetricSummary:
    """Mean, sample standard deviation, and fixed percentiles of one metric."""

    mean: float
    std: Optional[float]
    percentiles: Dict[int, float]


@dataclass
class GroupSummary:
    """Pooled spores for one photo, sample, probe, type, or the whole run."""

    level: str
    type_name: str
    probe_name: str
    sample_name: str
    image_name: str
    n_measured: int
    n_fallback: int
    metrics: Dict[str, MetricSummary]


def morphology_settings(
    config_manager,
    confidence: Optional[float] = None,
    count_inside_square: Optional[bool] = None,
) -> MorphologySettings:
    """Read morphology settings. A passed confidence or square mode replaces yaml."""
    chosen = (
        config_manager.get_float_param("morphology_confidence")
        if confidence is None
        else confidence
    )
    padding_ratio = config_manager.get_float_param("morphology_padding_ratio")
    montage_columns = config_manager.get_int_param("morphology_montage_columns")
    scale = _optional_scale(config_manager.get_param("microns_per_pixel"))
    if isinstance(chosen, bool) or not isinstance(chosen, (int, float)):
        raise ConfigurationError(
            f"config.yaml 'morphology_confidence' must be a number, got {chosen!r}"
        )
    chosen = float(chosen)
    if not 0.0 <= chosen <= 1.0:
        raise ConfigurationError(
            "config.yaml 'morphology_confidence' must be between 0 and 1"
        )
    if not math.isfinite(padding_ratio) or padding_ratio < 0:
        raise ConfigurationError(
            "config.yaml 'morphology_padding_ratio' must be a finite number >= 0"
        )
    if montage_columns < 1:
        raise ConfigurationError(
            "config.yaml 'morphology_montage_columns' must be an integer >= 1"
        )
    square_size = config_manager.get_int_param("analysis_square_size")
    if count_inside_square is None:
        count_inside_square = config_manager.get_bool_param("count_inside_square")
    return MorphologySettings(
        confidence=chosen,
        padding_ratio=float(padding_ratio),
        montage_columns=int(montage_columns),
        microns_per_pixel=scale,
        canny_threshold1=config_manager.get_int_param("canny_threshold1"),
        canny_threshold2=config_manager.get_int_param("canny_threshold2"),
        analysis_square_size=square_size,
        count_inside_square=bool(count_inside_square) and square_size > 0,
    )


def measurement_from_contour(contour: np.ndarray) -> Dict[str, float]:
    """Axes, area, and shape of one contour. fitEllipse axes stay full lengths."""
    points = np.asarray(contour)
    if len(points) < 5:
        raise ValueError("An ellipse fit needs at least 5 contour points")
    (center_x, center_y), (axis_a, axis_b), angle = cv2.fitEllipse(points)
    return _geometry(
        major=_major(axis_a, axis_b),
        minor=_minor(axis_a, axis_b),
        angle_deg=_major_axis_angle(axis_a, axis_b, angle),
        center_x=float(center_x),
        center_y=float(center_y),
        area=float(cv2.contourArea(points)),
        perimeter=float(cv2.arcLength(points, True)),
        hull_area=float(cv2.contourArea(cv2.convexHull(points))),
    )


def summarize_groups(
    measurements: Sequence[SporeMeasurement],
    scale: Optional[float],
) -> List[GroupSummary]:
    """Percentiles of each group come from its pooled contour measurements."""
    groups = [("run", "", "", "", "", list(measurements))]
    groups.extend(_buckets(measurements, "type", lambda row: (row.type_name,)))
    groups.extend(
        _buckets(measurements, "probe", lambda row: (row.type_name, row.probe_name))
    )
    groups.extend(
        _buckets(
            measurements,
            "sample",
            lambda row: (row.type_name, row.probe_name, row.sample_name),
        )
    )
    groups.extend(
        _buckets(measurements, "photo", lambda row: (row.image_path,))
    )
    summaries = []
    for level, type_name, probe_name, sample_name, image_name, rows in groups:
        if level != "run" and not any((type_name, probe_name, sample_name, image_name)):
            continue
        summaries.append(_summarize_group(
            level, type_name, probe_name, sample_name, image_name, rows, scale
        ))
    return summaries


def measure_detections(
    image: Image.Image,
    detections: Sequence,
    settings: MorphologySettings,
    *,
    image_path: str,
    type_name: str = "",
    probe_name: str = "",
    sample_name: str = "",
    class_names: Optional[Sequence[str]] = None,
    population: str = "yolo",
) -> List[SporeMeasurement]:
    """Measure accepted YOLO or SAHI boxes. The square filter is optional."""
    accepted_names = set(class_names or ["spore"])
    width, height = image.size
    measured: List[SporeMeasurement] = []
    for detection in detections:
        bbox, confidence, class_name = _detection_parts(detection)
        if class_name not in accepted_names:
            continue
        if confidence is None or confidence < settings.confidence:
            continue
        center_x = (bbox[0] + bbox[2]) / 2.0
        center_y = (bbox[1] + bbox[3]) / 2.0
        if not _keeps_center(center_x, center_y, width, height, settings):
            continue
        measured.append(
            _measure_box(
                image,
                bbox,
                settings,
                image_path=image_path,
                type_name=type_name,
                probe_name=probe_name,
                sample_name=sample_name,
                spore_index=len(measured) + 1,
                confidence=float(confidence),
                population=population,
            )
        )
    return measured


def measure_opencv_contours(
    image: Image.Image,
    contours: Sequence[np.ndarray],
    settings: MorphologySettings,
    *,
    image_path: str,
    type_name: str = "",
    probe_name: str = "",
    sample_name: str = "",
) -> List[SporeMeasurement]:
    """Measure contours the OpenCV detector already accepted."""
    width, height = image.size
    gray = np.array(image.convert("L"))
    measured: List[SporeMeasurement] = []
    for contour in contours:
        try:
            geometry = measurement_from_contour(contour)
        except (cv2.error, ValueError):
            continue
        if not _keeps_center(
            geometry["center_x"],
            geometry["center_y"],
            width,
            height,
            settings,
        ):
            continue
        bounds = cv2.boundingRect(np.asarray(contour))
        bbox = (bounds[0], bounds[1], bounds[0] + bounds[2], bounds[1] + bounds[3])
        crop, _origin, clipped = _padded_crop(image, bbox, settings.padding_ratio)
        measured.append(
            SporeMeasurement(
                image_path=image_path,
                type_name=type_name,
                probe_name=probe_name,
                sample_name=sample_name,
                spore_index=len(measured) + 1,
                confidence=None,
                center_x=geometry["center_x"],
                center_y=geometry["center_y"],
                major_px=geometry["major_px"],
                minor_px=geometry["minor_px"],
                area_px=geometry["area_px"],
                equivalent_diameter_px=geometry["equivalent_diameter_px"],
                axis_ratio=geometry["axis_ratio"],
                eccentricity=geometry["eccentricity"],
                circularity=geometry["circularity"],
                solidity=geometry["solidity"],
                angle_deg=geometry["angle_deg"],
                source=CONTOUR_SOURCE,
                edge_support_ratio=None,
                intensity_diff=_intensity_diff(gray, np.asarray(contour)),
                bbox_width_px=float(bounds[2]),
                bbox_height_px=float(bounds[3]),
                crop_clipped=clipped,
                population="opencv_accepted",
                crop=crop,
            )
        )
    return measured


def build_montage(
    measurements: Sequence[SporeMeasurement],
    columns: int,
    microns_per_pixel: Optional[float],
) -> Optional[np.ndarray]:
    """Stitch contour crops at their own size. Columns wrap onto the next row."""
    cells = [
        item
        for item in measurements
        if item.source == CONTOUR_SOURCE and item.crop is not None and item.crop.size
    ]
    if not cells:
        return None
    columns = max(1, int(columns))
    row_images = []
    for start in range(0, len(cells), columns):
        row_images.append(_montage_row(cells[start:start + columns], microns_per_pixel))
    width = max(image.shape[1] for image in row_images)
    height = sum(image.shape[0] for image in row_images)
    canvas = np.full((height, width, 3), 180, dtype=np.uint8)
    y = 0
    for image in row_images:
        canvas[y:y + image.shape[0], 0:image.shape[1]] = image
        y += image.shape[0]
    return canvas


def write_morphology_reports(
    measurements: Sequence[SporeMeasurement],
    output_dir: Path,
    settings: MorphologySettings,
) -> Dict[str, Path]:
    """Write the spore table, the pooled summary, and one montage per photo."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    scale = settings.microns_per_pixel
    reports = {
        "morphology_csv": _write_csv(measurements, output_dir / "morphology_spores.csv", scale),
        "morphology_excel": _write_excel(
            measurements, output_dir / "morphology_summary.xlsx", settings
        ),
        "morphology_markdown": _write_markdown(
            measurements, output_dir / "morphology_summary.md", settings
        ),
    }
    montage_dir = output_dir / "montages"
    by_image: Dict[str, List[SporeMeasurement]] = {}
    for item in measurements:
        by_image.setdefault(item.image_path, []).append(item)
    for image_path, rows in by_image.items():
        montage = build_montage(rows, settings.montage_columns, scale)
        if montage is None:
            continue
        montage_dir.mkdir(parents=True, exist_ok=True)
        destination = montage_dir / f"{Path(image_path).stem}_spores.jpg"
        cv2.imwrite(str(destination), cv2.cvtColor(montage, cv2.COLOR_RGB2BGR))
        reports[f"montage:{Path(image_path).stem}"] = destination
    return reports


def micron_values(measurement: SporeMeasurement, scale: float) -> Dict[str, float]:
    """Linear sizes scale by microns per pixel. Area scales by its square."""
    return {
        "major_um": measurement.major_px * scale,
        "minor_um": measurement.minor_px * scale,
        "equivalent_diameter_um": measurement.equivalent_diameter_px * scale,
        "area_um2": measurement.area_px * scale * scale,
    }


def _optional_scale(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ConfigurationError(
            f"config.yaml 'microns_per_pixel' must be a positive number or null, got {value!r}"
        )
    scale = float(value)
    if not math.isfinite(scale) or scale <= 0:
        raise ConfigurationError(
            "config.yaml 'microns_per_pixel' must be a positive number or null"
        )
    return scale


def _major(axis_a: float, axis_b: float) -> float:
    return float(max(axis_a, axis_b))


def _minor(axis_a: float, axis_b: float) -> float:
    return float(min(axis_a, axis_b))


def _major_axis_angle(axis_a: float, axis_b: float, angle: float) -> float:
    """OpenCV's angle follows the width axis. Report the major axis instead."""
    major_angle = float(angle) if axis_a >= axis_b else float(angle) + 90.0
    return major_angle % 180.0


def _geometry(
    *,
    major: float,
    minor: float,
    angle_deg: float,
    center_x: float,
    center_y: float,
    area: float,
    perimeter: float,
    hull_area: float,
) -> Dict[str, float]:
    major = max(float(major), 0.0)
    minor = max(float(minor), 0.0)
    axis_ratio = 0.0 if major <= 0 else minor / major
    eccentricity = math.sqrt(max(0.0, 1.0 - axis_ratio * axis_ratio))
    circularity = 0.0 if perimeter <= 0 else (4.0 * math.pi * area) / (perimeter * perimeter)
    solidity = 0.0 if hull_area <= 0 else area / hull_area
    equivalent = 0.0 if area <= 0 else 2.0 * math.sqrt(area / math.pi)
    return {
        "center_x": float(center_x),
        "center_y": float(center_y),
        "major_px": major,
        "minor_px": minor,
        "area_px": float(area),
        "equivalent_diameter_px": float(equivalent),
        "axis_ratio": float(axis_ratio),
        "eccentricity": float(eccentricity),
        "circularity": float(circularity),
        "solidity": float(solidity),
        "angle_deg": float(angle_deg) % 180.0,
    }


def _detection_parts(detection) -> Tuple[Tuple[float, float, float, float], Optional[float], str]:
    if isinstance(detection, dict):
        x1, y1, x2, y2 = detection["bbox"]
        return (
            (float(x1), float(y1), float(x2), float(y2)),
            None if detection.get("confidence") is None else float(detection["confidence"]),
            str(detection.get("class_name") or "spore"),
        )
    return (
        tuple(float(value) for value in detection.bbox),
        float(detection.confidence),
        str(detection.class_name),
    )


def _measure_box(
    image: Image.Image,
    bbox: Tuple[float, float, float, float],
    settings: MorphologySettings,
    **identity,
) -> SporeMeasurement:
    width = max(0.0, bbox[2] - bbox[0])
    height = max(0.0, bbox[3] - bbox[1])
    crop, origin, clipped = _padded_crop(image, bbox, settings.padding_ratio)
    geometry = None
    source = BBOX_SOURCE
    edge_support = None
    intensity = None
    if crop is not None:
        found = _contour_in_crop(crop, origin, bbox, settings)
        if found is not None:
            geometry, source, edge_support, intensity = found
    if geometry is None:
        geometry = _bbox_ellipse(bbox)
        source = BBOX_SOURCE
    return SporeMeasurement(
        center_x=geometry["center_x"],
        center_y=geometry["center_y"],
        major_px=geometry["major_px"],
        minor_px=geometry["minor_px"],
        area_px=geometry["area_px"],
        equivalent_diameter_px=geometry["equivalent_diameter_px"],
        axis_ratio=geometry["axis_ratio"],
        eccentricity=geometry["eccentricity"],
        circularity=geometry["circularity"],
        solidity=geometry["solidity"],
        angle_deg=geometry["angle_deg"],
        source=source,
        edge_support_ratio=edge_support,
        intensity_diff=intensity,
        bbox_width_px=width,
        bbox_height_px=height,
        crop_clipped=clipped,
        crop=crop if source == CONTOUR_SOURCE else None,
        **identity,
    )


def _padded_crop(image: Image.Image, bbox, padding_ratio: float):
    """Return the RGB crop, its origin on the full image, and whether padding was clipped."""
    image_width, image_height = image.size
    box_width = max(0.0, float(bbox[2]) - float(bbox[0]))
    box_height = max(0.0, float(bbox[3]) - float(bbox[1]))
    pad = padding_ratio * max(box_width, box_height)
    raw_x1 = float(bbox[0]) - pad
    raw_y1 = float(bbox[1]) - pad
    raw_x2 = float(bbox[2]) + pad
    raw_y2 = float(bbox[3]) + pad
    clipped = raw_x1 < 0 or raw_y1 < 0 or raw_x2 > image_width or raw_y2 > image_height
    x1 = max(0, int(math.floor(raw_x1)))
    y1 = max(0, int(math.floor(raw_y1)))
    x2 = min(image_width, int(math.ceil(raw_x2)))
    y2 = min(image_height, int(math.ceil(raw_y2)))
    if x2 - x1 < 3 or y2 - y1 < 3:
        return None, (x1, y1), clipped
    crop = np.array(image.crop((x1, y1, x2, y2)).convert("RGB"))
    return crop, (x1, y1), clipped


def _contour_in_crop(crop: np.ndarray, origin, bbox, settings: MorphologySettings):
    """Largest outer contour that contains the box center, or a rectangle fallback."""
    from bees.opencv.image_proc import EdgeDetector, ImagePreprocessor

    preprocessed = ImagePreprocessor().preprocess(Image.fromarray(crop))
    edges, edges_strong = EdgeDetector().detect_edges(
        preprocessed, settings.canny_threshold1, settings.canny_threshold2
    )
    cleaned = EdgeDetector().clean_edges(edges)
    contours, _hierarchy = cv2.findContours(cleaned, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)
    center = (
        ((bbox[0] + bbox[2]) / 2.0) - origin[0],
        ((bbox[1] + bbox[3]) / 2.0) - origin[1],
    )
    containing = []
    long_contours = []
    for contour in contours:
        area = float(cv2.contourArea(contour))
        if len(contour) >= 5 and area > 0:
            long_contours.append(contour)
            if cv2.pointPolygonTest(contour, center, False) >= 0:
                containing.append((area, contour))
    if containing:
        contour = max(containing, key=lambda item: item[0])[1]
        try:
            geometry = measurement_from_contour(contour)
        except (cv2.error, ValueError):
            geometry = None
        else:
            geometry["center_x"] += origin[0]
            geometry["center_y"] += origin[1]
            return (
                geometry,
                CONTOUR_SOURCE,
                _edge_support(edges_strong, contour),
                _intensity_diff(preprocessed, contour),
            )
    if long_contours:
        contour = max(long_contours, key=lambda item: float(cv2.contourArea(item)))
        geometry = _rect_geometry(contour)
        if geometry is not None:
            geometry["center_x"] += origin[0]
            geometry["center_y"] += origin[1]
            return geometry, RECT_SOURCE, None, None
    return None


def _rect_geometry(contour: np.ndarray) -> Optional[Dict[str, float]]:
    try:
        (center_x, center_y), (axis_a, axis_b), angle = cv2.minAreaRect(contour)
    except cv2.error:
        return None
    if max(axis_a, axis_b) <= 0:
        return None
    perimeter = float(cv2.arcLength(contour, True))
    hull_area = float(cv2.contourArea(cv2.convexHull(contour)))
    area = float(cv2.contourArea(contour))
    return _geometry(
        major=_major(axis_a, axis_b),
        minor=_minor(axis_a, axis_b),
        angle_deg=_major_axis_angle(axis_a, axis_b, angle),
        center_x=float(center_x),
        center_y=float(center_y),
        area=area if area > 0 else float(axis_a * axis_b),
        perimeter=perimeter,
        hull_area=hull_area,
    )


def _bbox_ellipse(bbox) -> Dict[str, float]:
    width = max(0.0, float(bbox[2]) - float(bbox[0]))
    height = max(0.0, float(bbox[3]) - float(bbox[1]))
    semi_major = _major(width, height) / 2.0
    semi_minor = _minor(width, height) / 2.0
    area = math.pi * semi_major * semi_minor
    perimeter = _ellipse_perimeter(semi_major, semi_minor)
    return _geometry(
        major=_major(width, height),
        minor=_minor(width, height),
        angle_deg=0.0,
        center_x=(float(bbox[0]) + float(bbox[2])) / 2.0,
        center_y=(float(bbox[1]) + float(bbox[3])) / 2.0,
        area=area,
        perimeter=perimeter,
        hull_area=area,
    )


def _ellipse_perimeter(semi_major: float, semi_minor: float) -> float:
    """Ramanujan's approximation. A circle falls back to 2·π·r."""
    if semi_major <= 0 and semi_minor <= 0:
        return 0.0
    if semi_major == semi_minor:
        return 2.0 * math.pi * semi_major
    total = semi_major + semi_minor
    difference = ((semi_major - semi_minor) / total) ** 2
    return math.pi * total * (1.0 + (3.0 * difference) / (10.0 + math.sqrt(4.0 - 3.0 * difference)))


def _edge_support(edges: np.ndarray, contour: np.ndarray) -> Optional[float]:
    try:
        perimeter = np.zeros_like(edges)
        cv2.drawContours(perimeter, [contour], -1, 255, 1)
        perimeter_pixels = int(cv2.countNonZero(perimeter))
        if perimeter_pixels == 0:
            return None
        overlap = int(cv2.countNonZero(cv2.bitwise_and(edges, perimeter)))
        return overlap / float(perimeter_pixels)
    except cv2.error:
        return None


def _intensity_diff(gray: np.ndarray, contour: np.ndarray) -> Optional[float]:
    try:
        mask = np.zeros(gray.shape[:2], dtype=np.uint8)
        cv2.drawContours(mask, [contour], -1, 255, -1)
        if int(cv2.countNonZero(mask)) == 0:
            return None
        inside = float(cv2.mean(gray, mask=mask)[0])
        return abs(inside - float(np.mean(gray)))
    except cv2.error:
        return None


def _buckets(measurements, level: str, key_of):
    buckets: Dict[tuple, List[SporeMeasurement]] = {}
    for item in measurements:
        key = key_of(item)
        if not any(key):
            continue
        buckets.setdefault(key, []).append(item)
    rows = []
    for key in sorted(buckets):
        type_name = key[0] if level in {"type", "probe", "sample"} else ""
        probe_name = key[1] if level in {"probe", "sample"} else ""
        sample_name = key[2] if level == "sample" else ""
        image_name = Path(key[0]).name if level == "photo" else ""
        if level == "type":
            probe_name = ""
            sample_name = ""
        rows.append((level, type_name, probe_name, sample_name, image_name, buckets[key]))
    return rows


def _summarize_group(level, type_name, probe_name, sample_name, image_name, rows, scale):
    measured = [item for item in rows if item.source == CONTOUR_SOURCE]
    metrics = {}
    if measured:
        for name in _summary_names(scale):
            values = np.array(_metric_values(measured, name, scale), dtype=float)
            std = float(np.std(values, ddof=1)) if len(values) >= 2 else None
            percentiles = {
                int(percent): float(value)
                for percent, value in zip(
                    PERCENTILES,
                    np.percentile(values, PERCENTILES, method="linear"),
                )
            }
            metrics[name] = MetricSummary(
                mean=float(np.mean(values)),
                std=std,
                percentiles=percentiles,
            )
    return GroupSummary(
        level=level,
        type_name=type_name,
        probe_name=probe_name,
        sample_name=sample_name,
        image_name=image_name,
        n_measured=len(measured),
        n_fallback=len(rows) - len(measured),
        metrics=metrics,
    )


def _summary_names(scale: Optional[float]) -> Iterable[str]:
    names = list(SUMMARY_METRICS)
    if scale is not None:
        names.extend(("major_um", "minor_um", "equivalent_diameter_um", "area_um2"))
    return names


def _metric_values(measurements, name: str, scale: Optional[float]) -> List[float]:
    if name == "major_um":
        return [item.major_px * scale for item in measurements]
    if name == "minor_um":
        return [item.minor_px * scale for item in measurements]
    if name == "equivalent_diameter_um":
        return [item.equivalent_diameter_px * scale for item in measurements]
    if name == "area_um2":
        return [item.area_px * scale * scale for item in measurements]
    return [float(getattr(item, name)) for item in measurements]


def _detail_fields(scale: Optional[float]):
    fields = list(DETAIL_FIELDS)
    if scale is not None:
        fields.extend(MICRON_FIELDS)
    return fields


def _write_csv(measurements, path: Path, scale: Optional[float]) -> Path:
    fields = _detail_fields(scale)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[key for key, _label in fields])
        writer.writeheader()
        for item in measurements:
            writer.writerow({
                key: _csv_value(item.table_values(scale).get(key))
                for key, _label in fields
            })
    return path


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.6f}"
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def _write_excel(measurements, path: Path, settings: MorphologySettings) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from bees.reporting import apply_header_style, apply_thin_borders

    workbook = Workbook()
    spores = workbook.active
    spores.title = "Споры"
    summary = workbook.create_sheet("Сводка")
    _write_spore_sheet(spores, measurements, settings, apply_header_style, apply_thin_borders, Font, Alignment)
    _write_summary_sheet(
        summary, measurements, settings, apply_header_style, apply_thin_borders, Font, Alignment, get_column_letter
    )
    workbook.save(path)
    return path


def _keeps_center(
    center_x: float,
    center_y: float,
    width: int,
    height: int,
    settings: MorphologySettings,
) -> bool:
    """The whole frame is measured unless the inside-square mode is on."""
    if not settings.count_inside_square or settings.analysis_square_size <= 0:
        return True
    return point_inside_analysis_square(
        center_x, center_y, width, height, settings.analysis_square_size
    )


def _population_lines(measurements, settings: MorphologySettings) -> List[str]:
    populations = sorted({item.population for item in measurements}) or ["yolo"]
    if settings.count_inside_square and settings.analysis_square_size > 0:
        scope = "Перцентили посчитаны по спорам внутри квадрата анализа."
    else:
        scope = "Перцентили посчитаны по спорам всего кадра."
    lines = [
        scope,
        "Источник: " + ", ".join(populations) + ".",
    ]
    if "opencv_accepted" in populations:
        lines.append("Форма этих спор уже прошла геометрические пороги детектора.")
    if "yolo" in populations:
        lines.append(f"Порог уверенности: {settings.confidence:.2f}.")
    if settings.microns_per_pixel is None:
        lines.append("Масштаб не задан.")
    else:
        lines.append(f"Масштаб: {settings.microns_per_pixel:g} мкм/пиксель.")
    return lines


def _write_spore_sheet(worksheet, measurements, settings, apply_header_style, apply_thin_borders, font, alignment):
    fields = _detail_fields(settings.microns_per_pixel)
    for line in _population_lines(measurements, settings):
        worksheet.append([line])
    worksheet.append([])
    from openpyxl.utils import get_column_letter

    worksheet.append([label for _key, label in fields])
    header_row = worksheet.max_row
    apply_header_style(
        [worksheet.cell(row=header_row, column=column) for column in range(1, len(fields) + 1)],
        wrap_text=True,
    )
    worksheet.row_dimensions[header_row].height = 36
    for column, (_key, label) in enumerate(fields, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = min(28, max(12, len(label) + 2))
    for item in measurements:
        values = item.table_values(settings.microns_per_pixel)
        worksheet.append([values[key] for key, _label in fields])
    last_row = max(header_row, worksheet.max_row)
    apply_thin_borders(worksheet, header_row, last_row, 1, len(fields))
    for cell in worksheet[1]:
        cell.font = font(bold=True)
        cell.alignment = alignment(wrap_text=True)


def _write_summary_sheet(
    worksheet, measurements, settings, apply_header_style, apply_thin_borders, font, alignment, get_column_letter
):
    for line in _population_lines(measurements, settings):
        worksheet.append([line])
    worksheet.append([])
    metric_names = list(_summary_names(settings.microns_per_pixel))
    identity = ["Уровень", "Тип", "Проба", "Образец", "Фото", "Измерено", "Запасные"]
    stat_labels = ["среднее", "ст. откл.", *[f"p{percent}" for percent in PERCENTILES]]
    top = list(identity)
    bottom = [""] * len(identity)
    for name in metric_names:
        top.append(METRIC_LABELS[name])
        top.extend([""] * (len(stat_labels) - 1))
        bottom.extend(stat_labels)
    worksheet.append(top)
    header_row = worksheet.max_row
    worksheet.append(bottom)
    for column in range(1, len(identity) + 1):
        worksheet.merge_cells(
            start_row=header_row,
            start_column=column,
            end_row=header_row + 1,
            end_column=column,
        )
    for index in range(len(metric_names)):
        start = len(identity) + 1 + index * len(stat_labels)
        worksheet.merge_cells(
            start_row=header_row,
            start_column=start,
            end_row=header_row,
            end_column=start + len(stat_labels) - 1,
        )
    column_count = len(top)
    apply_header_style(
        [
            worksheet.cell(row=row, column=column)
            for row in (header_row, header_row + 1)
            for column in range(1, column_count + 1)
        ],
        wrap_text=True,
    )
    worksheet.row_dimensions[header_row].height = 22
    worksheet.row_dimensions[header_row + 1].height = 18
    for group in summarize_groups(measurements, settings.microns_per_pixel):
        row = [
            LEVEL_LABELS[group.level],
            group.type_name,
            group.probe_name,
            group.sample_name,
            group.image_name,
            group.n_measured,
            group.n_fallback,
        ]
        for name in metric_names:
            metric = group.metrics.get(name)
            if metric is None:
                row.extend([None] * (2 + len(PERCENTILES)))
                continue
            row.append(metric.mean)
            row.append(metric.std)
            row.extend(metric.percentiles[percent] for percent in PERCENTILES)
        worksheet.append(row)
    last_row = max(header_row + 1, worksheet.max_row)
    apply_thin_borders(worksheet, header_row, last_row, 1, column_count)
    for column in range(1, column_count + 1):
        width = 14 if column <= len(identity) else 12
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for cell in worksheet[1]:
        cell.font = font(bold=True)
        cell.alignment = alignment(wrap_text=True)


def _write_markdown(measurements, path: Path, settings: MorphologySettings) -> Path:
    groups = summarize_groups(measurements, settings.microns_per_pixel)
    lines = ["# Морфометрия спор", ""]
    lines.extend(f"{line}" for line in _population_lines(measurements, settings))
    lines.append("")
    for group in groups:
        if group.level not in {"run", "probe"}:
            continue
        title = "Весь запуск" if group.level == "run" else group.probe_name or "Проба"
        if group.level == "probe" and group.type_name:
            title = f"{group.type_name} / {title}"
        lines.append(f"## {title}")
        lines.append("")
        lines.append(f"Измерено: {group.n_measured}. Запасные: {group.n_fallback}.")
        for name in ("major_px", "minor_px", "axis_ratio"):
            metric = group.metrics.get(name)
            if metric is None:
                continue
            low = metric.percentiles[5]
            high = metric.percentiles[95]
            lines.append(
                f"{METRIC_LABELS[name]}: медиана {metric.percentiles[50]:.4f}, "
                f"p5–p95 {low:.4f}–{high:.4f}."
            )
        if settings.microns_per_pixel is not None:
            for name in ("major_um", "minor_um"):
                metric = group.metrics.get(name)
                if metric is None:
                    continue
                lines.append(
                    f"{METRIC_LABELS[name]}: медиана {metric.percentiles[50]:.4f}, "
                    f"p5–p95 {metric.percentiles[5]:.4f}–{metric.percentiles[95]:.4f}."
                )
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _montage_row(cells: Sequence[SporeMeasurement], scale: Optional[float]) -> np.ndarray:
    tiles = [_label_tile(item, scale) for item in cells]
    height = max(tile.shape[0] for tile in tiles)
    width = sum(tile.shape[1] for tile in tiles) + MONTAGE_PAD * (len(tiles) + 1)
    row = np.full((height, width, 3), 180, dtype=np.uint8)
    x = MONTAGE_PAD
    for tile in tiles:
        row[0:tile.shape[0], x:x + tile.shape[1]] = tile
        x += tile.shape[1] + MONTAGE_PAD
    return row


def _label_tile(measurement: SporeMeasurement, scale: Optional[float]) -> np.ndarray:
    crop = measurement.crop
    if crop.ndim == 2:
        crop = cv2.cvtColor(crop, cv2.COLOR_GRAY2RGB)
    tile_height = crop.shape[0] + MONTAGE_LABEL_HEIGHT + MONTAGE_PAD
    tile_width = crop.shape[1]
    tile = np.full((tile_height, tile_width, 3), 180, dtype=np.uint8)
    tile[0:crop.shape[0], 0:crop.shape[1]] = crop
    confidence = "" if measurement.confidence is None else f" {measurement.confidence:.2f}"
    first = f"#{measurement.spore_index}{confidence} {measurement.major_px:.1f}x{measurement.minor_px:.1f} px"
    _draw_label(tile, first, crop.shape[0] + 12)
    if scale is not None:
        second = (
            f"{measurement.major_px * scale:.2f}x{measurement.minor_px * scale:.2f} um"
        )
        _draw_label(tile, second, crop.shape[0] + 26)
    return tile


def _draw_label(tile: np.ndarray, text: str, baseline: int) -> None:
    cv2.putText(
        tile,
        text,
        (2, baseline),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.35,
        (20, 20, 20),
        1,
        cv2.LINE_AA,
    )

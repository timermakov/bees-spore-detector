"""Morphology measurements, percentiles, and reports."""

import csv
import math
import sys
from pathlib import Path

import numpy as np
import pytest
from PIL import Image
from openpyxl import load_workbook

_REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_REPO_ROOT))

from bees.config_loader import ConfigurationError, create_config_manager
from bees.morphology import (
    BBOX_SOURCE,
    CONTOUR_SOURCE,
    SporeMeasurement,
    build_montage,
    measure_detections,
    measurement_from_contour,
    morphology_settings,
    summarize_groups,
    write_morphology_reports,
)
from bees.yolo.detector import Detection


def _ellipse_contour(cx: float, cy: float, semi_major: float, semi_minor: float) -> np.ndarray:
    angles = np.linspace(0, 2 * math.pi, 90, endpoint=False)
    points = np.stack(
        [cx + semi_major * np.cos(angles), cy + semi_minor * np.sin(angles)],
        axis=1,
    ).astype(np.float32)
    return points.reshape(-1, 1, 2)


def _settings(**overrides):
    values = {
        "confidence": 0.7,
        "padding_ratio": 0.1,
        "montage_columns": 2,
        "microns_per_pixel": None,
        "canny_threshold1": 40,
        "canny_threshold2": 120,
        "analysis_square_size": 80,
    }
    values.update(overrides)
    from bees.morphology import MorphologySettings

    return MorphologySettings(**values)


def _spore(**overrides) -> SporeMeasurement:
    values = {
        "image_path": "photo.png",
        "type_name": "Type_A",
        "probe_name": "Probe_1",
        "sample_name": "Sample_1",
        "spore_index": 1,
        "confidence": 0.9,
        "center_x": 50.0,
        "center_y": 50.0,
        "major_px": 10.0,
        "minor_px": 6.0,
        "area_px": 47.0,
        "equivalent_diameter_px": 7.7,
        "axis_ratio": 0.6,
        "eccentricity": 0.8,
        "circularity": 0.9,
        "solidity": 0.95,
        "angle_deg": 15.0,
        "source": CONTOUR_SOURCE,
        "edge_support_ratio": None,
        "intensity_diff": None,
        "bbox_width_px": 12.0,
        "bbox_height_px": 8.0,
        "crop_clipped": False,
        "population": "yolo",
        "crop": None,
    }
    values.update(overrides)
    return SporeMeasurement(**values)


def _yaml(directory: Path, body: str) -> Path:
    path = directory / "config.yaml"
    path.write_text(body, encoding="utf-8")
    return path


class TestContourGeometry:
    def test_known_ellipse_axes_area_and_eccentricity(self):
        # Arrange
        contour = _ellipse_contour(80, 70, semi_major=40, semi_minor=20)

        # Act
        geometry = measurement_from_contour(contour)

        # Assert
        assert geometry["major_px"] == pytest.approx(80, abs=1.5)
        assert geometry["minor_px"] == pytest.approx(40, abs=1.5)
        assert geometry["area_px"] == pytest.approx(math.pi * 40 * 20, rel=0.05)
        assert geometry["axis_ratio"] == pytest.approx(0.5, abs=0.03)
        assert geometry["eccentricity"] == pytest.approx(math.sqrt(0.75), abs=0.03)
        assert geometry["equivalent_diameter_px"] == pytest.approx(
            2 * math.sqrt(geometry["area_px"] / math.pi), rel=0.01
        )


class TestPercentiles:
    def test_linear_percentiles_and_sample_std(self):
        # Arrange
        rows = [_spore(major_px=float(value), spore_index=value) for value in range(1, 11)]

        # Act
        summary = summarize_groups(rows, scale=None)[0].metrics["major_px"]

        # Assert
        assert summary.mean == pytest.approx(5.5)
        assert summary.std == pytest.approx(math.sqrt(82.5 / 9))
        assert summary.percentiles[5] == pytest.approx(1.45)
        assert summary.percentiles[25] == pytest.approx(3.25)
        assert summary.percentiles[50] == pytest.approx(5.5)
        assert summary.percentiles[75] == pytest.approx(7.75)
        assert summary.percentiles[95] == pytest.approx(9.55)

    def test_fallback_rows_leave_the_median_unchanged(self):
        # Arrange
        rows = [
            _spore(major_px=10, spore_index=1),
            _spore(major_px=30, spore_index=2),
            _spore(major_px=1000, spore_index=3, source=BBOX_SOURCE),
        ]

        # Act
        group = summarize_groups(rows, scale=None)[0]

        # Assert
        assert group.n_measured == 2
        assert group.n_fallback == 1
        assert group.metrics["major_px"].percentiles[50] == pytest.approx(20)


class TestDetectionCrop:
    def test_low_confidence_and_outside_square_are_omitted(self):
        # Arrange
        image = Image.new("RGB", (100, 100), color="white")
        settings = _settings(analysis_square_size=40, confidence=0.7, count_inside_square=True)
        detections = [
            Detection((40, 40, 60, 60), confidence=0.9),
            Detection((40, 40, 60, 60), confidence=0.6),
            Detection((0, 0, 4, 4), confidence=0.95),
            Detection((40, 40, 60, 60), confidence=0.95, class_name="dust"),
        ]

        # Act
        measured = measure_detections(
            image,
            detections,
            settings,
            image_path="photo.png",
            class_names=["spore"],
        )

        # Assert
        assert len(measured) == 1
        assert measured[0].confidence == pytest.approx(0.9)
        assert measured[0].source == BBOX_SOURCE
        assert measured[0].major_px == pytest.approx(20)
        assert measured[0].area_px == pytest.approx(math.pi * 10 * 10)

    def test_outside_spore_stays_in_the_whole_frame(self):
        # Arrange
        image = Image.new("RGB", (100, 100), color="white")
        settings = _settings(analysis_square_size=40, confidence=0.7)
        detections = [
            Detection((40, 40, 60, 60), confidence=0.9),
            Detection((0, 0, 4, 4), confidence=0.95),
            Detection((40, 40, 60, 60), confidence=0.6),
        ]

        # Act
        measured = measure_detections(image, detections, settings, image_path="photo.png")

        # Assert
        assert [row.confidence for row in measured] == pytest.approx([0.9, 0.95])

    def test_drawn_ellipse_is_measured_from_its_contour(self):
        # Arrange
        canvas = np.full((160, 160, 3), 255, dtype=np.uint8)
        import cv2

        cv2.ellipse(canvas, (80, 80), (36, 16), 0, 0, 360, (0, 0, 0), -1)
        image = Image.fromarray(canvas)
        detection = Detection((40, 60, 120, 104), confidence=0.91)

        # Act
        measured = measure_detections(
            image,
            [detection],
            _settings(analysis_square_size=120, confidence=0.7, padding_ratio=0.2),
            image_path="ellipse.png",
        )

        # Assert
        assert measured[0].source == CONTOUR_SOURCE
        assert measured[0].major_px == pytest.approx(72, abs=8)
        assert measured[0].minor_px == pytest.approx(32, abs=8)
        assert measured[0].crop is not None
        assert measured[0].crop.shape[0] >= 32
        assert measured[0].crop.shape[1] >= 72


class TestReports:
    def test_null_scale_omits_micron_columns(self, tmp_path):
        # Arrange
        rows = [_spore(major_px=10, area_px=80), _spore(major_px=30, area_px=120, spore_index=2)]
        settings = _settings(microns_per_pixel=None)

        # Act
        reports = write_morphology_reports(rows, tmp_path, settings)

        # Assert
        with reports["morphology_csv"].open(encoding="utf-8") as handle:
            header = next(csv.reader(handle))
        sheet = load_workbook(reports["morphology_excel"])["Сводка"]
        text = reports["morphology_markdown"].read_text(encoding="utf-8")
        assert "major_um" not in header
        assert "area_um2" not in header
        assert any(cell.value and "Масштаб не задан" in str(cell.value) for row in sheet.iter_rows(max_row=6) for cell in row)
        assert "Большая ось, пикс." in text
        assert "всего кадра" in text
        assert "мкм" not in text

    def test_positive_scale_converts_axis_and_area(self, tmp_path):
        # Arrange
        rows = [_spore(major_px=10, minor_px=4, area_px=80, equivalent_diameter_px=8)]
        settings = _settings(microns_per_pixel=0.5)

        # Act
        reports = write_morphology_reports(rows, tmp_path, settings)

        # Assert
        with reports["morphology_csv"].open(encoding="utf-8") as handle:
            table = list(csv.DictReader(handle))
        assert "major_um" in table[0]
        assert float(table[0]["major_um"]) == pytest.approx(5)
        assert float(table[0]["minor_um"]) == pytest.approx(2)
        assert float(table[0]["area_um2"]) == pytest.approx(20)
        assert "Масштаб: 0.5 мкм/пиксель" in reports["morphology_markdown"].read_text(encoding="utf-8")

    def test_opencv_summary_names_the_detector_population(self, tmp_path):
        # Arrange
        rows = [_spore(population="opencv_accepted", confidence=None)]

        # Act
        reports = write_morphology_reports(rows, tmp_path, _settings())

        # Assert
        text = reports["morphology_markdown"].read_text(encoding="utf-8")
        assert "opencv_accepted" in text
        assert "геометрические пороги детектора" in text


class TestMontage:
    def test_crop_pixels_keep_their_size_and_columns_wrap(self):
        # Arrange
        first = np.zeros((5, 8, 3), dtype=np.uint8)
        first[0, 0] = (255, 0, 0)
        first[4, 7] = (0, 0, 255)
        second = np.full((6, 4, 3), (0, 255, 0), dtype=np.uint8)
        third = np.full((3, 3, 3), (255, 255, 0), dtype=np.uint8)
        rows = [
            _spore(crop=first, spore_index=1),
            _spore(crop=second, spore_index=2),
            _spore(crop=third, spore_index=3),
        ]

        # Act
        montage = build_montage(rows, columns=2, microns_per_pixel=None)

        # Assert
        red = np.argwhere(np.all(montage == (255, 0, 0), axis=2))[0]
        blue = np.argwhere(np.all(montage == (0, 0, 255), axis=2))[0]
        green = np.argwhere(np.all(montage == (0, 255, 0), axis=2))[0]
        yellow = np.argwhere(np.all(montage == (255, 255, 0), axis=2))[0]
        assert int(blue[0] - red[0]) == 4
        assert int(blue[1] - red[1]) == 7
        assert abs(int(green[0]) - int(red[0])) <= 2
        assert int(yellow[0]) > int(red[0]) + 5


class TestMorphologyConfig:
    def test_missing_confidence_is_named(self, tmp_path):
        # Arrange
        path = _yaml(
            tmp_path,
            "\n".join(
                [
                    "data_dir: data",
                    "results_dir: results",
                    "canny_threshold1: 20",
                    "canny_threshold2: 40",
                    "analysis_square_size: 30",
                    "morphology_padding_ratio: 0.1",
                    "morphology_montage_columns: 3",
                    "microns_per_pixel: null",
                ]
            )
            + "\n",
        )

        # Act
        with pytest.raises(ConfigurationError) as caught:
            morphology_settings(create_config_manager(path))

        # Assert
        assert "morphology_confidence" in str(caught.value)

    def test_repository_config_exposes_morphology_settings(self):
        # Arrange
        manager = create_config_manager(_REPO_ROOT / "config.yaml")

        # Act
        settings = morphology_settings(manager)

        # Assert
        assert 0.0 <= settings.confidence <= 1.0
        assert settings.padding_ratio >= 0
        assert settings.montage_columns >= 1
        assert settings.microns_per_pixel is None or settings.microns_per_pixel > 0
        assert settings.count_inside_square is False

    def test_missing_square_mode_is_named(self, tmp_path):
        # Arrange
        path = _yaml(
            tmp_path,
            "\n".join(
                [
                    "data_dir: data",
                    "results_dir: results",
                    "canny_threshold1: 20",
                    "canny_threshold2: 40",
                    "analysis_square_size: 30",
                    "morphology_confidence: 0.7",
                    "morphology_padding_ratio: 0.1",
                    "morphology_montage_columns: 3",
                    "microns_per_pixel: null",
                ]
            )
            + "\n",
        )

        # Act
        with pytest.raises(ConfigurationError) as caught:
            morphology_settings(create_config_manager(path))

        # Assert
        assert "count_inside_square" in str(caught.value)

    def test_zero_square_turns_the_inside_mode_off(self, tmp_path):
        # Arrange
        path = _yaml(
            tmp_path,
            "\n".join(
                [
                    "data_dir: data",
                    "results_dir: results",
                    "canny_threshold1: 20",
                    "canny_threshold2: 40",
                    "analysis_square_size: 0",
                    "count_inside_square: true",
                    "morphology_confidence: 0.7",
                    "morphology_padding_ratio: 0.1",
                    "morphology_montage_columns: 3",
                    "microns_per_pixel: null",
                ]
            )
            + "\n",
        )

        # Act
        settings = morphology_settings(create_config_manager(path))

        # Assert
        assert settings.analysis_square_size == 0
        assert settings.count_inside_square is False
